"""Builds apps/demo/public/ejemplo-ventas.xlsx, the sample workbook of the xlsx demo.

Fictional data that exercises what the viewer draws: merged title, frozen header, column
widths, dates, currency, percentages, negative numbers in red, booleans, wrapped and
overflowing text, borders, fills and several sheets (one with a hidden column).

    python3 scripts/make-sample-xlsx.py
"""
import datetime
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

random.seed(7)
OUT = os.path.join(os.path.dirname(__file__), "..", "apps", "demo", "public", "ejemplo-ventas.xlsx")

wb = Workbook()
ws = wb.active
ws.title = "Ventas"
thin = Side(style="thin", color="A6A6A6")
medium = Side(style="medium", color="1F4E78")

ws.merge_cells("A1:H1")
ws["A1"] = "Informe de ventas · primer semestre 2026"
ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30
ws["A2"] = "Datos ficticios para la demostración del visor. Los importes están en euros e incluyen descuentos por volumen."
ws["A2"].font = Font(italic=True, color="595959")

headers = ["Fecha", "Región", "Producto", "Unidades", "Precio", "Total", "Margen", "Pagado"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=3, column=i, value=h)
    c.font = Font(bold=True, color="1F4E78")
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.border = Border(bottom=medium)
    c.alignment = Alignment(horizontal="center")

regions = ["Andina", "Caribe", "Pacífico", "Orinoquía", "Amazonía"]
products = [("Teclado mecánico", 59.9), ("Ratón inalámbrico", 24.5), ("Monitor 27\"", 219.0), ("Base para portátil", 35.0), ("Auriculares", 79.99)]
day = datetime.date(2026, 1, 2)
row = 4
for n in range(60):
    product, price = random.choice(products)
    units = random.randint(1, 40)
    total = round(units * price * (0.9 if units >= 20 else 1.0), 2)
    margin = round(random.uniform(-0.08, 0.35), 4)
    values = [day, random.choice(regions), product, units, price, total, margin, random.random() > 0.25]
    for col, value in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=value)
        c.border = Border(bottom=thin)
    ws.cell(row=row, column=1).number_format = "dd/mm/yyyy"
    ws.cell(row=row, column=5).number_format = '#,##0.00 "€"'
    ws.cell(row=row, column=6).number_format = '#,##0.00 "€"'
    ws.cell(row=row, column=7).number_format = '0.0%;[Red]-0.0%'
    if n % 2:
        for col in range(1, 9):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor="F2F2F2")
    day += datetime.timedelta(days=random.randint(1, 4))
    row += 1

ws.cell(row=row, column=5, value="Total").font = Font(bold=True)
ws.cell(row=row, column=6, value=sum(ws.cell(row=r, column=6).value for r in range(4, row)))
ws.cell(row=row, column=6).number_format = '#,##0.00 "€"'
ws.cell(row=row, column=6).font = Font(bold=True)
ws.cell(row=row, column=6).border = Border(top=thin, bottom=Side(style="double"))

widths = {"A": 12, "B": 13, "C": 22, "D": 10, "E": 12, "F": 14, "G": 10, "H": 9}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "B4"

summary = wb.create_sheet("Resumen")
summary["A1"] = "Resumen por región"
summary["A1"].font = Font(size=14, bold=True)
summary.append([])
summary.append(["Región", "Notas internas", "Pedidos", "Importe", "Cumplimiento"])
for c in summary[3]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="E2EFDA")
    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, region in enumerate(regions):
    r = 4 + i
    summary.cell(row=r, column=1, value=region)
    summary.cell(row=r, column=2, value="columna oculta")
    summary.cell(row=r, column=3, value=random.randint(8, 20))
    summary.cell(row=r, column=4, value=round(random.uniform(4000, 15000), 2)).number_format = "#,##0.00"
    summary.cell(row=r, column=5, value=round(random.uniform(0.6, 1.2), 3)).number_format = "0%"
    for col in range(1, 6):
        summary.cell(row=r, column=col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
summary.column_dimensions["A"].width = 16
summary.column_dimensions["B"].hidden = True
summary.column_dimensions["D"].width = 14
summary.column_dimensions["E"].width = 14
summary.sheet_view.showGridLines = False

notes = wb.create_sheet("Notas")
notes.column_dimensions["A"].width = 60
notes["A1"] = "Notas"
notes["A1"].font = Font(size=14, bold=True)
notes["A2"] = (
    "Este libro se generó para probar el visor de @dimarborda/xlsx-editor. El texto largo de esta celda "
    "tiene activado el ajuste de texto, así que se reparte en varias líneas dentro del ancho de la columna."
)
notes["A2"].alignment = Alignment(wrap_text=True, vertical="top")
notes.row_dimensions[2].height = 60
notes["A4"] = "Centrado"
notes["A4"].alignment = Alignment(horizontal="center")
notes["A5"] = "A la derecha"
notes["A5"].alignment = Alignment(horizontal="right")
notes["A6"] = "Con sangría"
notes["A6"].alignment = Alignment(indent=2)
notes["A7"] = "Tachado y subrayado"
notes["A7"].font = Font(strike=True, underline="single", color="C00000")
notes["A8"] = 1234567.891
notes["A9"] = 0.000012345
notes["A10"] = datetime.datetime(2026, 10, 17, 18, 45)
notes["A10"].number_format = "dddd d \"de\" mmmm \"de\" yyyy, h:mm AM/PM"

wb.save(OUT)
print("guardado", os.path.normpath(OUT))
